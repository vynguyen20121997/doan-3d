# -*- coding: utf-8 -*-
"""
Sinh báo cáo Checkpoint 1 theo đúng format mẫu nhom_9-checkpoint1.docx.

Cách làm: mở file mẫu để kế thừa toàn bộ style (Heading có đánh số Chương/1.1/1.1.1,
Table Grid, List Bullet, khổ giấy, lề), xoá sạch phần thân rồi viết nội dung mới.

    python tools/build_checkpoint1.py
"""
import os
import copy
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
# Nguồn style: file mẫu của lớp; nếu không còn thì dùng chính bản đã sinh
# trước đó (nó kế thừa đủ style từ file mẫu) để script luôn chạy được.
MAU = os.path.join(os.path.expanduser("~"), "Downloads", "nhom_9-checkpoint1.docx")
OUT_DOCX = os.path.join(ROOT, "CHECKPOINT1-Cham-cong-dinh-vi-3D.docx")
DG = os.path.join(ROOT, "docs", "diagrams")
SHOT = os.path.join(ROOT, "docs")

MAX_W, MAX_H = 6.0, 7.6

# Danh sach nhom 1 - trich tu DS_IE402_F31_Danhsachlop
THANH_VIEN = [
    ("23730217", "Nguyễn Hữu Tín"),      # nhóm trưởng
    ("23730177", "Nguyễn Minh Khôi"),
    ("23730226", "Nguyễn Tường Vy"),
    ("23730204", "Võ Thành Nhân"),
]

# ---------------------------------------------------------------------------
if not os.path.exists(MAU):
    MAU = OUT_DOCX
doc = Document(MAU)

# numPr dùng để tắt đánh số cho các Heading 1 kiểu "DANH MỤC HÌNH"
_no_num_src = None
for p in doc.paragraphs:
    if p.text.strip() == "DANH MỤC HÌNH":
        pPr = p._p.find(qn("w:pPr"))
        if pPr is not None:
            n = pPr.find(qn("w:numPr"))
            if n is not None:
                _no_num_src = copy.deepcopy(n)
        break

# Xoá sạch thân tài liệu (giữ lại sectPr cuối để không mất khổ giấy/lề)
body = doc.element.body
for child in list(body):
    if child.tag != qn("w:sectPr"):
        body.remove(child)


# ---------------------------------------------------------------------------
#  Tiện ích
# ---------------------------------------------------------------------------
def para(text="", style=None, bold=False, italic=False, size=None,
         align=None, indent=None, space_after=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        if size:
            r.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    if indent is not None:
        p.paragraph_format.first_line_indent = Inches(indent)
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p


def body_text(text):
    """Đoạn văn thân bài: thụt dòng đầu, canh đều."""
    return para(text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=0.3)


# Google Docs khong ho tro danh so heading nhieu cap nhu Word, nen so thu tu
# duoc viet thang vao noi dung heading va tat danh so tu dong cua style.
_cnt = {"c1": 0, "c2": 0, "c3": 0}


def _tat_danh_so(p):
    if _no_num_src is None:
        return
    pPr = p._p.get_or_add_pPr()
    cu = pPr.find(qn("w:numPr"))
    if cu is not None:
        pPr.remove(cu)
    pPr.insert(0, copy.deepcopy(_no_num_src))


def h1(text, numbered=True):
    if numbered:
        _cnt["c1"] += 1
        _cnt["c2"] = 0
        _cnt["c3"] = 0
        text = "Chương %d: %s" % (_cnt["c1"], text)
    p = doc.add_paragraph(text, style="Heading 1")
    _tat_danh_so(p)
    return p


def h2(text):
    _cnt["c2"] += 1
    _cnt["c3"] = 0
    p = doc.add_paragraph("%d.%d	%s" % (_cnt["c1"], _cnt["c2"], text), style="Heading 2")
    _tat_danh_so(p)
    return p


def h3(text):
    _cnt["c3"] += 1
    p = doc.add_paragraph("%d.%d.%d	%s" % (_cnt["c1"], _cnt["c2"], _cnt["c3"], text),
                          style="Heading 3")
    _tat_danh_so(p)
    return p


def bullets(items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def mono(text):
    """Khối sơ đồ/mã dạng chữ đều — dùng cho kiến trúc, pipeline, công thức."""
    for line in text.split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(0.35)
        r = p.add_run(line if line else " ")
        r.font.name = "Consolas"
        r.font.size = Pt(9)
    para()


def table(headers, rows, widths=None, caption=None):
    if caption:
        c = para(caption, align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
        c.runs[0].italic = True
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, hcell in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(hcell)
        r.bold = True
        r.font.size = Pt(10)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(10)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    para()
    return t


def image(filename, caption=None, folder=DG):
    path = filename if os.path.isabs(filename) else os.path.join(folder, filename)
    if not os.path.exists(path):
        para("[thiếu ảnh: %s]" % os.path.basename(path), italic=True)
        return
    w_px, h_px = Image.open(path).size
    w_in = MAX_W
    h_in = w_in * h_px / w_px
    if h_in > MAX_H:
        h_in = MAX_H
        w_in = h_in * w_px / h_px
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(w_in))
    if caption:
        c = para(caption, align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
        c.runs[0].italic = True
    para()


def page_break():
    doc.add_page_break()


def toc_field():
    """Chèn field MỤC LỤC — mở Word bấm Ctrl+A rồi F9 để sinh nội dung."""
    p = doc.add_paragraph()
    r = p.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), r'TOC \o "1-3" \h \z \u')
    inner = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "(Bấm Ctrl+A rồi F9 trong Word để sinh mục lục)"
    inner.append(t)
    fld.append(inner)
    r._r.addnext(fld)


C = WD_ALIGN_PARAGRAPH.CENTER

# ===========================================================================
#  TRANG BÌA
# ===========================================================================
para("ĐẠI HỌC QUỐC GIA THÀNH PHỐ HỒ CHÍ MINH", align=C, bold=True, size=13)
para("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN", align=C, bold=True, size=13)
para("TRUNG TÂM PHÁT TRIỂN CNTT", align=C, bold=True, size=13)
para(); para(); para()
para("BÁO CÁO ĐỒ ÁN MÔN", align=C, bold=True, size=15)
para("HỆ THỐNG THÔNG TIN ĐỊA LÝ 3 CHIỀU", align=C, bold=True, size=15)
para()
para("ĐỀ TÀI:", align=C, bold=True, size=13)
para("XÂY DỰNG HỆ THỐNG CHẤM CÔNG ĐỊNH VỊ 3D", align=C, bold=True, size=17)
para("THEO MÔ HÌNH KHỐI KHÔNG GIAN PHÂN TẦNG", align=C, bold=True, size=17)
para("CHO TOÀ NHÀ VĂN PHÒNG", align=C, bold=True, size=17)
para(); para(); para()
para("GVHD: ThS. PHAN THANH VŨ", align=C, bold=True, size=13)
para("Nhóm 1", align=C, bold=True, size=13)
para("Sinh viên thực hiện:", align=C, bold=True, size=13)
for _mssv, _ten in THANH_VIEN:
    para("%s	%s" % (_mssv, _ten), align=C, size=13)
para(); para(); para()
para("TP. Hồ Chí Minh, tháng 9 năm 2026", align=C, italic=True, size=13)
page_break()

# ===========================================================================
#  TÓM TẮT
# ===========================================================================
para("TÓM TẮT ĐỀ TÀI", align=C, bold=True, size=14)
para()
body_text(
    "Đề tài xây dựng một hệ thống chấm công định vị dựa trên mô hình khối không gian "
    "ba chiều, nhằm khắc phục hạn chế căn bản của các ứng dụng chấm công hiện nay: "
    "hàng rào địa lý hai chiều không phân biệt được người lao động đang ở tầng nào "
    "trong một toà nhà văn phòng cao tầng.")
body_text(
    "Hệ thống đề xuất mô hình MAP — Khối chấm công phân tầng (Multi-floor Attendance "
    "Prism). Mỗi văn phòng được biểu diễn bằng một khối lăng trụ có đa giác nền là "
    "footprint phần sàn thuê, giới hạn dưới và giới hạn trên là cao độ tuyệt đối của "
    "sàn và trần, kèm hai biên dung sai bù cho sai số định vị. Phép xác thực chấm "
    "công vì vậy chuyển từ bài toán “điểm thuộc đa giác” sang bài toán “điểm thuộc "
    "khối”, và sinh ra một trạng thái mới mà hệ thống hai chiều không thể tạo ra: "
    "đúng toà nhà nhưng sai tầng.")
body_text(
    "Giải pháp kỹ thuật sử dụng ArcGIS Maps SDK for JavaScript để hiển thị và dựng "
    "khối ba chiều trên nền địa hình, PostgreSQL kết hợp PostGIS để lưu trữ hình học "
    "có toạ độ cao độ và thực hiện truy vấn không gian, Node.js cho tầng dịch vụ và "
    "React cho giao diện. Dữ liệu nền công trình được trích xuất từ OpenStreetMap qua "
    "Overpass API. Bản dựng thử nghiệm trên toà nhà IFC One Saigon đã chứng minh hệ "
    "thống phát hiện được tình huống chấm công sai tầng mà hàng rào hai chiều bỏ sót.")
page_break()

# ===========================================================================
#  MỤC LỤC / DANH MỤC
# ===========================================================================
para("MỤC LỤC", align=C, bold=True, size=14)
toc_field()
page_break()

h1("DANH MỤC HÌNH", numbered=False)
for line in [
    "Hình 2.1\tSơ đồ quan hệ thực thể (ERD) của hệ thống",
    "Hình 2.2\tSơ đồ trường hợp sử dụng",
    "Hình 2.3\tSơ đồ luồng dữ liệu mức ngữ cảnh",
    "Hình 2.4\tSơ đồ luồng dữ liệu mức 1",
    "Hình 2.5\tSơ đồ trình tự chấm công vào",
    "Hình 2.6\tSơ đồ trình tự dựng khối MAP",
    "Hình 2.7\tSơ đồ trình tự duyệt đơn giải trình",
    "Hình 3.1\tBa khối MAP trong cùng một toà nhà (bản dựng thử nghiệm)",
    "Hình 3.2	Màn hình chấm công của hệ thống đã cài đặt",
    "Hình 3.3	Bảng điều khiển của quản lý",
    "Hình 4.1\tKết quả tình huống chấm công sai tầng",
]:
    para(line)

h1("DANH MỤC BẢNG", numbered=False)
for line in [
    "Bảng 1.1\tĐối tượng sử dụng và quyền chính",
    "Bảng 2.1\tCác lớp dữ liệu không gian",
    "Bảng 2.2\tThành phần của khối MAP",
    "Bảng 2.3\tSo sánh các mô hình biểu diễn đối tượng ba chiều",
    "Bảng 2.4\tBa trạng thái kết quả xác thực",
    "Bảng 2.5\tCông nghệ sử dụng",
    "Bảng 3.1\tBốn quy tắc phát hiện bất thường",
    "Bảng 4.1\tKết quả ba tình huống kiểm thử",
    "Bảng 3.2	Các điểm cuối API của hệ thống",
    "Bảng 4.2	Kết quả chạy thật trên PostgreSQL + PostGIS",
    "Bảng 4.3	Kết quả kiểm thử đầu cuối trên hệ thống đã cài đặt",
    "Bảng B.1\tKế hoạch triển khai theo giai đoạn",
]:
    para(line)
page_break()

# ===========================================================================
#  CHƯƠNG 1
# ===========================================================================
h1("TỔNG QUAN ĐỀ TÀI")

h2("Giới thiệu")
body_text(
    "Chấm công định vị (geolocation-based attendance) là hình thức xác nhận sự có mặt "
    "của người lao động bằng toạ độ thiết bị di động, thay cho máy chấm công vân tay "
    "hoặc thẻ từ đặt cố định. Hình thức này trở nên phổ biến khi doanh nghiệp có nhiều "
    "chi nhánh, nhân viên làm việc tại công trường hoặc di chuyển liên tục, khiến việc "
    "yêu cầu tất cả về một máy chấm công vật lý trở nên không khả thi.")
body_text(
    "Hầu hết ứng dụng chấm công định vị hiện nay đều dựa trên hàng rào địa lý hai "
    "chiều (2D geofence): một hình tròn bán kính R quanh một toạ độ, hoặc một đa giác "
    "phẳng. Hệ thống chỉ trả lời được câu hỏi “người này có đứng trong vùng hay không” "
    "trên mặt phẳng, hoàn toàn bỏ qua chiều cao.")
body_text(
    "Đề tài xây dựng hệ thống chấm công định vị sử dụng mô hình khối ba chiều, cho "
    "phép phân biệt vị trí người lao động theo cả chiều đứng, áp dụng cho bối cảnh "
    "toà nhà văn phòng cao tầng tại Thành phố Hồ Chí Minh. Hệ thống đồng thời cung cấp "
    "công cụ trực quan hoá để người quản lý đối chiếu khi có tranh chấp về công, và "
    "bộ quy tắc phát hiện các bản ghi chấm công bất thường.")

h2("Lý do chọn đề tài")
body_text(
    "Hàng rào hai chiều thất bại trong chính bối cảnh mà chấm công định vị được dùng "
    "nhiều nhất — toà nhà văn phòng cao tầng ở đô thị. Bốn vấn đề cụ thể như sau.")
body_text(
    "Thứ nhất, hệ thống không phân biệt được theo chiều cao. Một cao ốc bốn mươi tầng "
    "có hàng chục doanh nghiệp thuê các tầng khác nhau; nhân viên công ty A ở tầng 5 "
    "và nhân viên công ty B ở tầng 22 có cùng một toạ độ kinh độ – vĩ độ, nên hàng rào "
    "hai chiều chấp nhận cả hai như nhau. Nghiêm trọng hơn, một người ngồi ở quán cà "
    "phê tầng trệt hoặc đang ở hầm giữ xe vẫn được ghi nhận là chấm công hợp lệ cho "
    "văn phòng tầng 22.")
body_text(
    "Thứ hai, sai số định vị trong đô thị rất lớn do hiệu ứng hẻm núi đô thị (urban "
    "canyon): tín hiệu vệ tinh bị các khối nhà cao tầng che chắn và phản xạ nhiều "
    "đường, khiến sai số ngang có thể lên tới hàng chục mét — đủ để một người đứng ở "
    "toà nhà kế bên vẫn lọt vào hàng rào. Trên mô hình hai chiều không có cách nào "
    "phân tích hay khoanh vùng hiện tượng này; phải có mô hình khối của các công trình "
    "xung quanh mới mô tả được vùng bị che khuất.")
body_text(
    "Thứ ba, hệ thống không phát hiện được gian lận theo chiều đứng. Các thủ thuật phổ "
    "biến như dùng ứng dụng giả lập vị trí hoặc nhờ đồng nghiệp chấm hộ ngay dưới sảnh "
    "đều tạo ra bản ghi đúng toạ độ nhưng sai độ cao hoặc sai quy luật di chuyển; dữ "
    "liệu hai chiều không lưu giữ thông tin cần thiết để phát hiện.")
body_text(
    "Thứ tư, hệ thống không trực quan hoá được để đối chiếu. Khi phát sinh tranh chấp "
    "về công, người quản lý cần nhìn thấy nhân viên đã chấm công ở đâu; một điểm trên "
    "bản đồ phẳng không cho biết điểm đó nằm trong hay ngoài khối văn phòng.")
body_text(
    "Bốn vấn đề trên đều có gốc chung là thiếu chiều thứ ba trong mô hình dữ liệu. "
    "Đây chính là lý do đề tài được lựa chọn cho môn Hệ thống thông tin địa lý 3 chiều: "
    "bài toán chỉ giải được khi bổ sung trục cao độ, chứ không phải bổ sung để trang trí.")

h2("Mục tiêu đề tài")
h3("Mục tiêu tổng quát")
body_text(
    "Xây dựng hệ thống chấm công định vị dựa trên mô hình khối không gian ba chiều "
    "phân tầng, cho phép xác thực vị trí người lao động theo cả phương ngang lẫn "
    "phương đứng, phát hiện các bản ghi bất thường và trực quan hoá kết quả trên nền "
    "bản đồ ba chiều.")

h3("Mục tiêu cụ thể")
bullets([
    "Xây dựng mô hình dữ liệu MAP — khối chấm công phân tầng — và định nghĩa hình thức "
    "các thành phần của khối.",
    "Cài đặt thuật toán kiểm tra bao hàm khối có xét dung sai, trả về ba trạng thái "
    "thay vì kết quả nhị phân đúng/sai.",
    "Xây dựng cơ sở dữ liệu không gian lưu trữ toà nhà, văn phòng, khối chấm công, "
    "nhân sự, ca làm việc và bản ghi chấm công.",
    "Cài đặt bốn quy tắc phát hiện bất thường: sai tầng, dịch chuyển bất khả thi, độ "
    "chính xác bất thường và trùng thiết bị.",
    "Dựng giao diện bản đồ ba chiều hiển thị khối toà nhà, khối văn phòng và điểm chấm "
    "công đặt đúng cao độ.",
    "Cung cấp công cụ cho quản trị viên dựng khối vùng chấm công trực tiếp trên bản đồ "
    "ba chiều.",
    "Xây dựng luồng nghiệp vụ đầy đủ: chấm công, giải trình, duyệt và tổng hợp báo cáo công.",
])

h2("Đối tượng sử dụng")
table(
    ["Đối tượng", "Quyền chính"],
    [
        ["Nhân viên",
         "Chấm công vào/ra; xem lịch sử chấm công của bản thân; xem vị trí đã chấm "
         "trên bản đồ 3D; gửi đơn giải trình."],
        ["Quản lý",
         "Toàn bộ quyền của nhân viên; xem bảng điều khiển của phòng ban; xem danh "
         "sách cảnh báo bất thường; duyệt hoặc từ chối đơn giải trình; xuất báo cáo công."],
        ["Quản trị viên",
         "Quản lý toà nhà, văn phòng và khối vùng chấm công 3D; quản lý nhân viên, "
         "phòng ban, ca làm việc và thiết bị; cấu hình dung sai và ngưỡng cảnh báo."],
    ],
    widths=[1.3, 4.7],
    caption="Bảng 1.1 — Đối tượng sử dụng và quyền chính")

h2("Phạm vi đề tài")
body_text(
    "Việc số hoá toàn bộ nội thất một cao ốc với độ chính xác BIM đòi hỏi hồ sơ thiết "
    "kế và thiết bị khảo sát chuyên dụng, nằm ngoài khả năng của một đồ án môn học. "
    "Do đó đề tài tập trung xây dựng một hệ thống mẫu (prototype) với phạm vi như sau.")
bullets([
    "Một toà nhà văn phòng thử nghiệm, được dựng ở mức chi tiết LoD1 — khối đùn từ "
    "footprint và chiều cao.",
    "Ba văn phòng thuộc ba doanh nghiệp khác nhau, đặt ở ba dải tầng khác nhau trong "
    "cùng toà nhà, nhằm minh chứng trực tiếp cho hạn chế của hàng rào hai chiều.",
    "Dữ liệu footprint lấy từ OpenStreetMap; cao độ nền và chiều cao tầng lấy từ khảo "
    "sát và tài liệu công bố.",
    "Chức năng chấm công triển khai trên trình duyệt, sử dụng HTML5 Geolocation API; "
    "không phát triển ứng dụng di động gốc.",
    "Định vị trong nhà bằng Wi-Fi hoặc BLE beacon chỉ được nêu ở phần hướng phát triển, "
    "không cài đặt trong phạm vi đồ án.",
    "Chức năng tính lương và tích hợp với hệ thống nhân sự bên ngoài nằm ngoài phạm vi.",
])
page_break()

# ===========================================================================
#  CHƯƠNG 2
# ===========================================================================
h1("PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG")

h2("Dữ liệu của hệ thống")

h3("Dữ liệu không gian")
body_text(
    "Dữ liệu không gian là nền tảng để xác định quan hệ vị trí giữa người lao động và "
    "văn phòng. Hệ thống tổ chức dữ liệu theo các lớp sau.")
table(
    ["Lớp", "Kiểu hình học", "Nội dung"],
    [
        ["Building", "Polygon + chiều cao", "Footprint và khối bao của toà nhà văn phòng."],
        ["Office", "Polygon + dải tầng", "Phần sàn mà một doanh nghiệp thuê."],
        ["AttendanceZone", "PolygonZ (khối MAP)", "Khối không gian dùng để xác thực chấm công."],
        ["CheckinPoint", "PointZ", "Vị trí ba chiều của từng lần chấm công."],
        ["BuildingContext", "Polygon", "Các công trình lân cận, dùng phân tích che khuất tín hiệu."],
        ["Terrain", "Raster / DEM", "Địa hình nền, xác định cao độ mặt đất."],
    ],
    widths=[1.4, 1.6, 3.0],
    caption="Bảng 2.1 — Các lớp dữ liệu không gian")

h3("Dữ liệu toà nhà và cao độ")
body_text(
    "Mỗi toà nhà lưu bốn thông tin quyết định việc suy ra cao độ của các tầng: đa giác "
    "nền, cao độ mặt nền so với mực nước biển, tổng số tầng và chiều cao trung bình "
    "một tầng. Quy trình xử lý dữ liệu như sau.")
mono("OpenStreetMap (Overpass API)\n"
     "        |\n"
     "        v\n"
     "  Footprint (GeoJSON)  +  Khảo sát số tầng, chiều cao\n"
     "        |\n"
     "        v\n"
     "  PostGIS: bảng toa_nha (footprint, cao_do_nen, so_tang, chieu_cao_tang)")
body_text(
    "Khảo sát trên phạm vi Thành phố Hồ Chí Minh cho thấy trong 105.609 công trình có "
    "footprint trên OpenStreetMap, chỉ 3.944 công trình (khoảng 3,7 %) có thuộc tính "
    "chiều cao hoặc số tầng. Vì vậy dữ liệu độ cao của toà nhà thử nghiệm phải được "
    "nhập và kiểm chứng thủ công.")

h3("Khối vùng chấm công — mô hình MAP")
body_text(
    "Mô hình lăng trụ thuần tuý được biến đổi thành mô hình phục vụ riêng bài toán "
    "chấm công, đặt tên là MAP — Khối chấm công phân tầng (Multi-floor Attendance Prism).")
mono("MAP = ( P, z_min, z_max, dxy, dz )")
table(
    ["Thành phần", "Ý nghĩa"],
    [
        ["P", "Đa giác nền của phần sàn mà văn phòng thuê, hệ toạ độ WGS84."],
        ["z_min", "Cao độ tuyệt đối của sàn tầng thấp nhất văn phòng chiếm dụng."],
        ["z_max", "Cao độ tuyệt đối của trần tầng cao nhất."],
        ["dxy", "Biên dung sai ngang, bù cho sai số định vị đô thị."],
        ["dz", "Biên dung sai đứng, bù cho sai số đo cao độ."],
    ],
    widths=[1.2, 4.8],
    caption="Bảng 2.2 — Thành phần của khối MAP")
body_text("Cao độ được suy ra từ thông tin toà nhà và dải tầng văn phòng thuê:")
mono("z_min = z_nen + (tang_bat_dau - 1) x h_tang\n"
     "z_max = z_nen +  tang_ket_thuc      x h_tang")
body_text("Khối dùng để kiểm tra là khối MAP đã nới biên dung sai:")
mono("MAP+ = ( buffer(P, dxy),  z_min - dz,  z_max + dz )")

h3("Dữ liệu định vị thu từ thiết bị")
bullets([
    "Kinh độ, vĩ độ và cao độ do HTML5 Geolocation API trả về.",
    "Độ chính xác ngang và độ chính xác đứng, dùng để đánh giá mức tin cậy của bản ghi.",
    "Nguồn cao độ: đo được từ thiết bị, hoặc suy từ tầng do người dùng khai báo.",
    "Định danh thiết bị, phục vụ quy tắc phát hiện trùng thiết bị.",
    "Thời điểm chấm công, phục vụ quy tắc phát hiện dịch chuyển bất khả thi.",
])

h3("Dữ liệu nhân sự và ca làm việc")
bullets([
    "Công ty, phòng ban và danh sách nhân viên.",
    "Tài khoản người dùng và vai trò (nhân viên, quản lý, quản trị viên).",
    "Ca làm việc: giờ vào, giờ ra và số phút cho phép đi trễ.",
    "Phân ca theo ngày cho từng nhân viên.",
])

h3("Dữ liệu chấm công và cảnh báo")
bullets([
    "Bản ghi chấm công kèm trạng thái xác thực.",
    "Cảnh báo bất thường sinh từ bốn quy tắc R1 đến R4, kèm mức độ nghiêm trọng.",
    "Đơn giải trình của nhân viên và quyết định duyệt của quản lý.",
])

h2("Cơ sở lý thuyết và lựa chọn mô hình ba chiều")
body_text(
    "Các mô hình biểu diễn đối tượng ba chiều thường dùng trong hệ thống thông tin địa "
    "lý được so sánh trong bảng dưới đây.")
table(
    ["Mô hình", "Mô tả", "Ưu điểm", "Nhược điểm"],
    [
        ["2.5D Prism", "Đa giác nền được đùn lên theo chiều cao",
         "Nhẹ; dựng trực tiếp từ footprint và số tầng; kiểm tra bao hàm rất rẻ",
         "Không mô tả được phần nhô ra, mái vòm hay nội thất"],
        ["B-Rep", "Mô tả bằng tập mặt – cạnh – đỉnh",
         "Chính xác, biểu diễn được hình học phức tạp",
         "Dữ liệu nặng, khó dựng và khó truy vấn không gian"],
        ["CSG", "Tổ hợp Boolean các khối cơ bản",
         "Gọn với hình dạng quy tắc",
         "Không phù hợp với dữ liệu đo đạc thực tế"],
        ["Voxel / Octree", "Chia không gian thành các ô lập phương",
         "Tốt cho truy vấn thể tích và mô phỏng lan truyền",
         "Tốn bộ nhớ, mất độ chính xác ở biên"],
        ["TIN", "Lưới tam giác bất quy tắc",
         "Phù hợp mô tả địa hình",
         "Chỉ là bề mặt, không phải khối đặc"],
    ],
    widths=[1.0, 1.4, 1.9, 1.7],
    caption="Bảng 2.3 — So sánh các mô hình biểu diễn đối tượng ba chiều")
body_text(
    "Đề tài lựa chọn mô hình 2.5D Prism vì bốn lý do. Một là bài toán chỉ cần trả lời "
    "câu hỏi điểm có nằm trong khối văn phòng hay không, không cần mô tả chi tiết kiến "
    "trúc, nên lăng trụ là mức chi tiết vừa đủ, tương đương LoD1 trong chuẩn CityGML. "
    "Hai là phép kiểm tra tách được thành hai bước có chi phí thấp: kiểm tra điểm "
    "thuộc đa giác trên mặt phẳng bằng thuật toán ray casting với độ phức tạp tuyến "
    "tính theo số đỉnh, và so sánh cao độ với dải giới hạn. Ba là dữ liệu đầu vào sẵn "
    "có, không cần quét laser hay mô hình BIM. Bốn là mô hình được cả ArcGIS Maps SDK "
    "và PostGIS hỗ trợ trực tiếp.")
body_text(
    "Nhược điểm chấp nhận được của lựa chọn này là mọi tầng của một văn phòng bị coi "
    "là một khối liền, không mô tả được vách ngăn bên trong. Với bài toán chấm công, "
    "vốn chỉ cần xác định người lao động thuộc văn phòng nào, đây không phải hạn chế "
    "thực chất.")
body_text(
    "Phép kiểm tra bao hàm đối với bản ghi chấm công tại điểm Q gồm hai điều kiện:")
mono("trong_khoi(Q, MAP+)  <=>  pointInPolygon(Q.xy, buffer(P, dxy))\n"
     "                      AND  (z_min - dz) <= Q.alt <= (z_max + dz)")
body_text("Kết quả được phân thành ba trạng thái thay vì nhị phân đúng/sai.")
table(
    ["Trạng thái", "Điều kiện", "Ý nghĩa nghiệp vụ"],
    [
        ["HOP_LE", "Thoả cả điều kiện ngang và điều kiện đứng; độ chính xác đạt ngưỡng",
         "Ghi nhận công bình thường"],
        ["NGHI_NGO", "Thoả điều kiện ngang nhưng sai điều kiện đứng",
         "Đúng toà nhà, sai tầng — cần xác minh"],
        ["NGOAI_VUNG", "Không thoả điều kiện ngang",
         "Từ chối, đề nghị gửi đơn giải trình"],
    ],
    widths=[1.1, 2.8, 2.1],
    caption="Bảng 2.4 — Ba trạng thái kết quả xác thực")
body_text(
    "Trạng thái NGHI_NGO chính là giá trị mà mô hình hai chiều không thể sinh ra. Đây "
    "là đóng góp cốt lõi của mô hình MAP.")

h2("Kiến trúc hệ thống đề xuất")
body_text(
    "Hệ thống được thiết kế theo kiến trúc nhiều tầng, tách biệt giao diện, dịch vụ "
    "xác thực không gian, dữ liệu nghiệp vụ và tài nguyên bản đồ.")
mono("CLIENT\n"
     "  |- Web Application (React / TypeScript)\n"
     "  |- 3D Viewer (ArcGIS Maps SDK for JS - SceneView)\n"
     "  |- Geolocation API (lat, lon, alt, accuracy)\n"
     "            |\n"
     "         REST API\n"
     "            |\n"
     "   Node.js / Express\n"
     "     |- Dich vu xac thuc khoi (kiem tra bao ham MAP)\n"
     "     |- Bo phat hien bat thuong (R1..R4)\n"
     "     |- Dich vu bao cao cong\n"
     "            |\n"
     "   PostgreSQL + PostGIS\n"
     "     |- Hinh hoc co toa do Z (PolygonZ, PointZ)\n"
     "     |- Du lieu nhan su, ca lam viec, cham cong\n"
     "            |\n"
     "   Nguon du lieu ngoai\n"
     "     |- OpenStreetMap / Overpass API (footprint)\n"
     "     |- ArcGIS basemap + world-elevation (nen, dia hinh)")
body_text(
    "ArcGIS SceneView đảm nhiệm việc hiển thị khối ba chiều trên nền địa hình thực. "
    "PostGIS lưu trữ hình học có toạ độ cao độ và thực hiện các truy vấn không gian; "
    "hàm ST_DWithin trên kiểu geography cho phép tính khoảng cách trực tiếp bằng mét, "
    "tránh việc phải chuyển đổi hệ toạ độ thủ công.")

h2("Công nghệ sử dụng")
table(
    ["Thành phần", "Công nghệ đề xuất"],
    [
        ["Bản đồ 3D", "ArcGIS Maps SDK for JavaScript 4.29"],
        ["Ký hiệu khối", "PolygonSymbol3D + ExtrudeSymbol3DLayer + SolidEdges3D"],
        ["Công cụ dựng khối", "SketchViewModel"],
        ["Frontend", "React + TypeScript + Vite"],
        ["Định vị", "HTML5 Geolocation API"],
        ["Backend", "Node.js + Express"],
        ["Database", "PostgreSQL"],
        ["Spatial Database", "PostGIS (PolygonZ, PointZ, ST_DWithin, ST_Distance)"],
        ["Dữ liệu nền", "OpenStreetMap qua Overpass API (giấy phép ODbL)"],
        ["Địa hình", "ArcGIS world-elevation"],
        ["Định dạng trao đổi", "GeoJSON"],
        ["Authentication", "JWT + bcrypt"],
        ["Triển khai", "Vercel (frontend) + Render (backend, PostGIS)"],
    ],
    widths=[1.8, 4.2],
    caption="Bảng 2.5 — Công nghệ sử dụng")
body_text(
    "Lưu ý kỹ thuật bắt buộc: Geolocation API chỉ hoạt động trên kết nối HTTPS hoặc "
    "localhost. Khi trình diễn trên điện thoại phải dùng tên miền có chứng chỉ hoặc "
    "công cụ tạo đường hầm; mở bằng địa chỉ IP nội bộ sẽ bị trình duyệt từ chối cấp "
    "quyền định vị.")

h2("Thiết kế cơ sở dữ liệu")
body_text(
    "Cơ sở dữ liệu liên kết dữ liệu không gian, dữ liệu nhân sự và dữ liệu nghiệp vụ "
    "chấm công. Các bảng chính gồm:")

_db = [
    ("CongTy", [("ma_cong_ty", "serial, khoá chính"), ("ten_cong_ty", "varchar"),
                ("ma_so_thue", "varchar, duy nhất")]),
    ("PhongBan", [("ma_phong_ban", "serial, khoá chính"), ("ma_cong_ty", "integer, khoá ngoại"),
                  ("ten_phong_ban", "varchar")]),
    ("ToaNha", [("ma_toa_nha", "serial, khoá chính"), ("ten", "varchar"), ("dia_chi", "varchar"),
                ("footprint", "GEOMETRY(Polygon, 4326)"), ("cao_do_nen", "numeric(7,2) — mét"),
                ("so_tang", "integer"), ("chieu_cao_tang", "numeric(5,2) — mét"),
                ("nguon_du_lieu", "varchar")]),
    ("VanPhong", [("ma_van_phong", "serial, khoá chính"), ("ma_cong_ty", "integer, khoá ngoại"),
                  ("ma_toa_nha", "integer, khoá ngoại"), ("ten", "varchar"),
                  ("tang_bat_dau", "integer"), ("tang_ket_thuc", "integer")]),
    ("KhuVucChamCong", [("ma_khu_vuc", "serial, khoá chính"),
                        ("ma_van_phong", "integer, khoá ngoại, duy nhất"),
                        ("da_giac_nen", "GEOMETRY(Polygon, 4326)"),
                        ("z_min", "numeric(7,2)"), ("z_max", "numeric(7,2)"),
                        ("dung_sai_ngang", "numeric(6,2) — mét"),
                        ("dung_sai_dung", "numeric(6,2) — mét"),
                        ("dang_hieu_luc", "boolean")]),
    ("NhanVien", [("ma_nhan_vien", "serial, khoá chính"), ("ma_cong_ty", "integer, khoá ngoại"),
                  ("ma_phong_ban", "integer, khoá ngoại"), ("ho_ten", "varchar"),
                  ("email", "varchar, duy nhất"), ("ngay_vao_lam", "date"),
                  ("dang_lam_viec", "boolean")]),
    ("VaiTro", [("ma_vai_tro", "serial, khoá chính"),
                ("ten_vai_tro", "varchar — NHAN_VIEN | QUAN_LY | QUAN_TRI")]),
    ("NguoiDung", [("ma_nguoi_dung", "serial, khoá chính"),
                   ("ma_nhan_vien", "integer, khoá ngoại, duy nhất"),
                   ("ma_vai_tro", "integer, khoá ngoại"), ("ten_dang_nhap", "varchar"),
                   ("mat_khau_bam", "varchar — bcrypt")]),
    ("ThietBi", [("ma_thiet_bi", "serial, khoá chính"), ("ma_nhan_vien", "integer, khoá ngoại"),
                 ("dinh_danh_thiet_bi", "varchar"), ("he_dieu_hanh", "varchar"),
                 ("duoc_tin_cay", "boolean")]),
    ("CaLamViec", [("ma_ca", "serial, khoá chính"), ("ten_ca", "varchar"),
                   ("gio_vao", "time"), ("gio_ra", "time"), ("tre_toi_da_phut", "integer")]),
    ("PhanCa", [("ma_phan_ca", "serial, khoá chính"), ("ma_nhan_vien", "integer, khoá ngoại"),
                ("ma_ca", "integer, khoá ngoại"), ("ngay_lam_viec", "date")]),
    ("BanGhiChamCong", [("ma_ban_ghi", "bigserial, khoá chính"),
                        ("ma_nhan_vien", "integer, khoá ngoại"),
                        ("ma_khu_vuc", "integer, khoá ngoại"),
                        ("ma_thiet_bi", "integer, khoá ngoại"), ("thoi_diem", "timestamptz"),
                        ("loai", "varchar — VAO | RA"),
                        ("vi_tri", "GEOMETRY(PointZ, 4326)"),
                        ("do_chinh_xac_ngang", "numeric(7,2)"),
                        ("do_chinh_xac_dung", "numeric(7,2)"),
                        ("nguon_cao_do", "varchar — THIET_BI | KHAI_BAO"),
                        ("tang_khai_bao", "integer"),
                        ("trang_thai", "varchar — HOP_LE | NGHI_NGO | NGOAI_VUNG")]),
    ("CanhBaoBatThuong", [("ma_canh_bao", "serial, khoá chính"),
                          ("ma_ban_ghi", "bigint, khoá ngoại"),
                          ("ma_quy_tac", "varchar — R1 | R2 | R3 | R4"),
                          ("muc_do", "varchar — THAP | TRUNG_BINH | CAO"),
                          ("mo_ta", "text"), ("da_xu_ly", "boolean")]),
    ("DonGiaiTrinh", [("ma_don", "serial, khoá chính"), ("ma_nhan_vien", "integer, khoá ngoại"),
                      ("ma_ban_ghi", "bigint, khoá ngoại, duy nhất"), ("ly_do", "text"),
                      ("trang_thai", "varchar — CHO_DUYET | DA_DUYET | TU_CHOI"),
                      ("nguoi_duyet", "integer, khoá ngoại"), ("thoi_diem_gui", "timestamptz")]),
]
for ten, cols in _db:
    h3(ten)
    table(["Thuộc tính", "Kiểu"], [[a, b] for a, b in cols], widths=[2.2, 3.8])

body_text(
    "Hai thuộc tính z_min và z_max của bảng KhuVucChamCong là thuộc tính dẫn xuất từ "
    "cao độ nền và dải tầng. Hệ thống vẫn lưu trữ chúng (phi chuẩn hoá có chủ đích) "
    "kèm một trigger tự suy giá trị khi để trống, nhằm tránh phải tính lại trong các "
    "truy vấn xác thực vốn được gọi rất thường xuyên.")
body_text("Các ràng buộc toàn vẹn cần bảo đảm:")
bullets([
    "tang_bat_dau nhỏ hơn hoặc bằng tang_ket_thuc, và tang_ket_thuc không vượt quá "
    "so_tang của toà nhà.",
    "z_min luôn nhỏ hơn z_max.",
    "Mỗi nhân viên chỉ có tối đa một bản ghi loại VAO chưa có bản ghi RA tương ứng trong ngày.",
    "Cặp (ma_nhan_vien, ngay_lam_viec) trong bảng PhanCa là khoá dự tuyển — một người "
    "chỉ được xếp một ca mỗi ngày.",
])

h2("Mô hình ERD tổng quát")
mono("cong_ty        1 - N  phong_ban\n"
     "cong_ty        1 - N  van_phong\n"
     "cong_ty        1 - N  nhan_vien\n"
     "toa_nha        1 - N  van_phong\n"
     "van_phong      1 - 1  khu_vuc_cham_cong\n"
     "phong_ban      1 - N  nhan_vien\n"
     "nhan_vien      1 - 1  nguoi_dung\n"
     "vai_tro        1 - N  nguoi_dung\n"
     "nhan_vien      1 - N  thiet_bi\n"
     "nhan_vien      1 - N  phan_ca\n"
     "ca_lam_viec    1 - N  phan_ca\n"
     "nhan_vien      1 - N  ban_ghi_cham_cong\n"
     "khu_vuc        1 - N  ban_ghi_cham_cong\n"
     "thiet_bi       1 - N  ban_ghi_cham_cong\n"
     "ban_ghi        1 - N  canh_bao_bat_thuong\n"
     "ban_ghi        1 - 0..1 don_giai_trinh")
image("01-erd.png", "Hình 2.1 — Sơ đồ quan hệ thực thể (ERD) của hệ thống")
body_text(
    "Ba thực thể mang tính không gian là điểm khác biệt so với một hệ thống chấm công "
    "thông thường: ToaNha giữ khối bao và là gốc để suy ra cao độ mọi tầng; "
    "KhuVucChamCong là hiện thực trực tiếp của khối MAP và có quan hệ một–một với "
    "văn phòng; BanGhiChamCong lưu điểm ba chiều do thiết bị báo về và là đầu vào của "
    "phép kiểm tra bao hàm.")

h2("Phân tích Use Case")
image("02-usecase.png", "Hình 2.2 — Sơ đồ trường hợp sử dụng")

h3("Nhân viên")
bullets([
    "Đăng nhập hệ thống.",
    "Chấm công vào và chấm công ra bằng vị trí thiết bị.",
    "Xác nhận tầng đang làm việc khi thiết bị không đo được cao độ.",
    "Xem lịch sử chấm công của bản thân.",
    "Xem vị trí đã chấm công trên bản đồ ba chiều.",
    "Gửi đơn giải trình cho bản ghi bị từ chối hoặc bị đánh dấu nghi ngờ.",
])

h3("Quản lý")
bullets([
    "Toàn bộ chức năng của nhân viên.",
    "Xem bảng điều khiển công của phòng ban.",
    "Xem danh sách cảnh báo bất thường theo bốn quy tắc.",
    "Xem lại vị trí ba chiều của bản ghi bị cảnh báo.",
    "Duyệt hoặc từ chối đơn giải trình.",
    "Xuất báo cáo công theo tháng.",
])

h3("Quản trị viên")
bullets([
    "Quản lý danh mục toà nhà: footprint, cao độ nền, số tầng, chiều cao tầng.",
    "Quản lý văn phòng và dải tầng thuê.",
    "Dựng khối vùng chấm công ba chiều trực tiếp trên bản đồ.",
    "Cấu hình dung sai ngang, dung sai đứng và ngưỡng cảnh báo.",
    "Quản lý nhân viên, phòng ban, ca làm việc và thiết bị.",
])

h3("Đặc tả UC — Chấm công vào/ra")
table(
    ["Mục", "Nội dung"],
    [
        ["Tác nhân", "Nhân viên"],
        ["Tiền điều kiện",
         "Đã đăng nhập; thiết bị đã cấp quyền định vị; nhân viên có ca làm việc trong ngày"],
        ["Luồng chính",
         "1. Nhân viên mở màn hình chấm công.\n"
         "2. Hệ thống yêu cầu vị trí từ thiết bị.\n"
         "3. Thiết bị trả về kinh độ, vĩ độ, cao độ và độ chính xác.\n"
         "4. Nhân viên xác nhận tầng đang làm việc.\n"
         "5. Hệ thống thực hiện kiểm tra bao hàm khối.\n"
         "6. Hệ thống áp bốn quy tắc phát hiện bất thường.\n"
         "7. Hệ thống lưu bản ghi kèm trạng thái và hiển thị kết quả."],
        ["Luồng thay thế 3a",
         "Thiết bị không trả được cao độ; hệ thống dùng tầng khai báo ở bước 4 và đánh "
         "dấu nguồn cao độ là KHAI_BAO."],
        ["Luồng ngoại lệ 5a",
         "Ngoài vùng ngang; hệ thống từ chối, hiển thị khoảng cách tới khối gần nhất và "
         "gợi ý gửi đơn giải trình."],
        ["Luồng ngoại lệ 5b",
         "Trong vùng ngang nhưng lệch tầng; hệ thống lưu trạng thái NGHI_NGO, sinh cảnh "
         "báo R1 và thông báo cho quản lý."],
        ["Hậu điều kiện", "Có một bản ghi chấm công mới với trạng thái được xác định."],
    ],
    widths=[1.3, 4.7])

h3("Đặc tả UC — Dựng khối vùng chấm công ba chiều")
table(
    ["Mục", "Nội dung"],
    [
        ["Tác nhân", "Quản trị viên"],
        ["Tiền điều kiện", "Toà nhà đã tồn tại trong hệ thống"],
        ["Luồng chính",
         "1. Chọn toà nhà trên bản đồ ba chiều.\n"
         "2. Vẽ đa giác nền bằng công cụ phác thảo, hoặc kế thừa footprint của toà nhà.\n"
         "3. Nhập dải tầng văn phòng thuê.\n"
         "4. Hệ thống tự tính z_min và z_max từ cao độ nền và chiều cao tầng.\n"
         "5. Nhập dung sai ngang và dung sai đứng.\n"
         "6. Hệ thống dựng khối xem trước.\n"
         "7. Lưu khối."],
        ["Hậu điều kiện", "Khối MAP có hiệu lực, dùng để xác thực các lần chấm công sau."],
    ],
    widths=[1.3, 4.7])

h2("Data Flow Diagram")
h3("DFD mức ngữ cảnh")
image("03-dfd0.png", "Hình 2.3 — Sơ đồ luồng dữ liệu mức ngữ cảnh")

h3("DFD mức 1")
image("04-dfd1.png", "Hình 2.4 — Sơ đồ luồng dữ liệu mức 1")
table(
    ["Tiến trình", "Đầu vào", "Xử lý", "Đầu ra"],
    [
        ["1.0 Xác thực người dùng", "Thông tin đăng nhập",
         "Đối chiếu kho người dùng, sinh token kèm vai trò", "Phiên làm việc"],
        ["2.0 Thu nhận vị trí 3D", "Yêu cầu chấm công, dữ liệu định vị",
         "Chuẩn hoá toạ độ, xác định nguồn cao độ", "Điểm 3D kèm độ chính xác"],
        ["3.0 Kiểm tra bao hàm khối", "Điểm 3D, khối MAP",
         "Ray casting trên mặt phẳng và so sánh dải cao độ", "Bản ghi kèm trạng thái"],
        ["4.0 Phát hiện bất thường", "Bản ghi mới, lịch sử gần nhất",
         "Áp bốn quy tắc R1 đến R4", "Cảnh báo kèm mức độ"],
        ["5.0 Xử lý giải trình", "Đơn của nhân viên, quyết định của quản lý",
         "Cập nhật trạng thái đơn và bản ghi liên quan", "Đơn đã duyệt"],
        ["6.0 Tổng hợp báo cáo", "Bản ghi chấm công, ca làm việc",
         "Tính giờ công, số lần đi muộn, số lần bất thường", "Báo cáo tháng"],
        ["7.0 Quản trị khối MAP", "Footprint, dải tầng, dung sai",
         "Suy ra z_min và z_max, kiểm tra ràng buộc", "Khối MAP"],
    ],
    widths=[1.3, 1.5, 1.9, 1.3])

h2("Sequence Diagram")
h3("Sequence chấm công vào")
image("05-sd1-cham-cong.png", "Hình 2.5 — Sơ đồ trình tự chấm công vào")
body_text(
    "Nhánh thứ hai của khối alt mô tả đúng tình huống mà hàng rào hai chiều bỏ sót: "
    "toạ độ ngang nằm trong vùng nhưng cao độ lệch khỏi dải tầng, hệ thống trả về "
    "trạng thái NGHI_NGO và sinh cảnh báo R1 thay vì chấp nhận bản ghi.")

h3("Sequence dựng khối vùng chấm công")
image("06-sd2-dung-khoi.png", "Hình 2.6 — Sơ đồ trình tự dựng khối MAP")

h3("Sequence duyệt đơn giải trình")
image("07-sd3-giai-trinh.png", "Hình 2.7 — Sơ đồ trình tự duyệt đơn giải trình")
page_break()

# ===========================================================================
#  CHƯƠNG 3
# ===========================================================================
h1("XÂY DỰNG HỆ THỐNG")

h2("Quy trình xây dựng dữ liệu không gian")
mono("OpenStreetMap / Overpass API\n"
     "        |\n"
     "        v\n"
     "  Footprint toa nha (OSM JSON)\n"
     "        |\n"
     "        v\n"
     "  Chuyen doi sang GeoJSON, chuan hoa thuoc tinh\n"
     "        |\n"
     "        v\n"
     "  Khao sat so tang / chieu cao / cao do nen\n"
     "        |\n"
     "        v\n"
     "  Nap vao PostGIS (bang toa_nha, van_phong)\n"
     "        |\n"
     "        v\n"
     "  Sinh khoi MAP (bang khu_vuc_cham_cong)\n"
     "        |\n"
     "        v\n"
     "  ArcGIS SceneView: dung khoi va hien thi")

h2("Chuẩn hoá hệ toạ độ")
body_text(
    "Toàn bộ dữ liệu hình học được lưu ở hệ toạ độ địa lý WGS84 (EPSG:4326) để thống "
    "nhất với dữ liệu do Geolocation API trả về và với dữ liệu OpenStreetMap. Vì "
    "EPSG:4326 dùng đơn vị độ, mọi phép tính khoảng cách đều phải thực hiện trên kiểu "
    "geography của PostGIS để nhận kết quả bằng mét, thay vì tính trực tiếp trên toạ "
    "độ độ. Ở phía trình duyệt, khoảng cách được tính bằng phép chiếu phẳng cục bộ với "
    "hệ số quy đổi phụ thuộc vĩ độ; ở quy mô một toà nhà, sai số của phép xấp xỉ này "
    "là không đáng kể.")

h2("Dựng khối toà nhà")
body_text(
    "Khối toà nhà được dựng bằng cách đùn đa giác footprint từ cao độ nền lên một "
    "chiều cao bằng tích của số tầng và chiều cao trung bình mỗi tầng. Lớp đồ hoạ chứa "
    "khối được đặt ở chế độ cao độ tuyệt đối để khối bám đúng vào cao độ thực thay vì "
    "bám theo bề mặt địa hình. Khối toà nhà được tô bán trong suốt nhằm cho phép nhìn "
    "xuyên vào các khối văn phòng bên trong.")

h2("Xây dựng khối MAP và suy cao độ")
body_text(
    "Mỗi văn phòng sinh ra một khối MAP kế thừa đa giác nền của toà nhà (hoặc một đa "
    "giác nhỏ hơn nếu văn phòng chỉ thuê một phần sàn). Giới hạn dưới và giới hạn trên "
    "được suy ra tự động từ dải tầng theo công thức đã nêu ở mục 2.1.3. Ví dụ với toà "
    "nhà có cao độ nền 5 mét và chiều cao tầng 4,65 mét, văn phòng thuê từ tầng 5 đến "
    "tầng 8 sẽ có giới hạn dưới 23,60 mét và giới hạn trên 42,20 mét.")

h2("Thuật toán kiểm tra bao hàm khối")
body_text(
    "Thuật toán gồm hai bước độc lập. Bước thứ nhất kiểm tra điểm thuộc đa giác trên "
    "mặt phẳng bằng ray casting; nếu điểm nằm ngoài, hệ thống tính khoảng cách ngắn "
    "nhất từ điểm tới các cạnh của đa giác và so với dung sai ngang. Bước thứ hai so "
    "sánh cao độ của điểm với dải giới hạn đã nới biên dung sai đứng. Kết quả tổ hợp "
    "của hai bước cho ra một trong ba trạng thái ở Bảng 2.4.")
mono("kiem_tra_bao_ham(diem, khoi, dxy, dz):\n"
     "    d = khoang_cach_toi_da_giac(diem.xy, khoi.P)   # met, = 0 neu nam trong\n"
     "    trong_ngang = (d <= dxy)\n"
     "    trong_dung  = (khoi.z_min - dz) <= diem.z <= (khoi.z_max + dz)\n"
     "\n"
     "    neu khong trong_ngang     -> NGOAI_VUNG\n"
     "    neu trong_dung            -> HOP_LE\n"
     "    nguoc lai                 -> NGHI_NGO")
body_text(
    "Trên PostGIS, bước thứ nhất được cài đặt bằng hàm ST_DWithin trên kiểu geography, "
    "bước thứ hai bằng phép so sánh trên giá trị ST_Z của điểm. Toàn bộ logic được gói "
    "trong một hàm PL/pgSQL để tầng ứng dụng chỉ cần gọi một lần cho mỗi lần chấm công.")

h2("Quy tắc phát hiện bất thường")
table(
    ["Mã", "Quy tắc", "Cơ sở phát hiện", "Mức độ"],
    [
        ["R1", "Sai tầng",
         "Thoả điều kiện ngang nhưng lệch cao độ vượt dung sai đứng", "Cao"],
        ["R2", "Dịch chuyển bất khả thi",
         "Vận tốc suy ra giữa hai bản ghi liên tiếp vượt ngưỡng cho phép", "Cao"],
        ["R3", "Độ chính xác bất thường",
         "Giá trị độ chính xác quá nhỏ so với thực tế đô thị (dấu hiệu giả lập vị trí) "
         "hoặc quá lớn để kết luận", "Trung bình"],
        ["R4", "Trùng thiết bị",
         "Nhiều nhân viên chấm công từ cùng một định danh thiết bị trong khoảng thời "
         "gian ngắn", "Trung bình"],
    ],
    widths=[0.5, 1.5, 3.2, 0.8],
    caption="Bảng 3.1 — Bốn quy tắc phát hiện bất thường")

h2("Thiết kế giao diện")
h3("Màn hình chấm công")
bullets([
    "Nút chấm công vào và chấm công ra cỡ lớn, tối ưu cho thao tác một tay trên điện thoại.",
    "Hiển thị độ chính xác định vị hiện tại để người dùng biết khi nào tín hiệu kém.",
    "Ô chọn tầng đang làm việc, dùng khi thiết bị không đo được cao độ.",
    "Bản đồ thu nhỏ hiển thị vị trí hiện tại so với khối văn phòng.",
    "Thông báo kết quả phân biệt rõ ba trạng thái bằng màu sắc và nội dung khác nhau.",
])

h3("Màn hình bản đồ ba chiều")
bullets([
    "Khối toà nhà bán trong suốt làm bối cảnh.",
    "Các khối MAP tô màu riêng theo từng doanh nghiệp, có viền khối rõ ràng.",
    "Điểm chấm công hiển thị dưới dạng hình cầu đặt đúng cao độ, màu theo trạng thái.",
    "Bộ lọc theo ngày, theo nhân viên và theo trạng thái.",
    "Thanh công cụ chuyển đổi giữa chế độ xem ba chiều và hai chiều.",
])
image(os.path.join(SHOT, "01-ba-khoi-MAP.png"),
      "Hình 3.1 — Ba khối MAP trong cùng một toà nhà: Alpha Tech (cam, tầng 5–8), "
      "Beta Finance (xanh, tầng 20–24), Gamma Media (tím, tầng 35–40)", folder=SHOT)

h2("Dashboard quản trị")
bullets([
    "Thống kê số lượt chấm công hợp lệ, nghi ngờ và ngoài vùng theo ngày và theo phòng ban.",
    "Danh sách cảnh báo bất thường, sắp xếp theo mức độ nghiêm trọng.",
    "Danh sách đơn giải trình chờ duyệt, kèm liên kết mở vị trí ba chiều của bản ghi.",
    "Bảng tổng hợp giờ công theo tháng, hỗ trợ xuất tệp.",
    "Công cụ quản trị khối MAP: thêm, sửa, vô hiệu hoá khối và điều chỉnh dung sai.",
])

h2("Cài đặt hệ thống")
body_text(
    "Hệ thống đã được cài đặt hoàn chỉnh chứ không dừng ở bản mô phỏng thuật toán. "
    "Máy chủ viết bằng Node.js với Express, nói chuyện trực tiếp với PostgreSQL/PostGIS "
    "và phục vụ luôn giao diện web tĩnh, nên toàn bộ hệ thống chạy trên một cổng duy "
    "nhất — thuận lợi cho việc trình diễn vì Geolocation API chấp nhận localhost.")
body_text(
    "Điểm đáng chú ý về mặt kiến trúc: phép kiểm tra bao hàm khối không được viết lại "
    "ở tầng ứng dụng mà gọi thẳng hàm kiem_tra_bao_ham() và kiem_tra_r2() đã cài đặt "
    "trong cơ sở dữ liệu. Nhờ vậy logic không gian chỉ tồn tại một bản duy nhất, đặt "
    "ngay cạnh dữ liệu, tránh tình trạng hai nơi cùng cài đặt rồi lệch nhau.")
body_text(
    "Mỗi lần chấm công được xử lý trong một giao dịch: ghi bản ghi, áp bốn quy tắc "
    "R1–R4 rồi ghi các cảnh báo tương ứng. Nếu bất kỳ bước nào lỗi thì toàn bộ được "
    "huỷ, không để lại bản ghi dở dang.")
table(
    ["Thành phần", "Công nghệ", "Vai trò"],
    [
        ["server/", "Node.js + Express + pg",
         "11 điểm cuối API, xác thực JWT, phân quyền ba vai trò"],
        ["web/", "HTML/CSS/JavaScript thuần + ArcGIS SDK",
         "Đăng nhập, chấm công, lịch sử, bảng điều khiển, bản đồ ba chiều"],
        ["db/", "PostgreSQL 17.6 + PostGIS 3.6.2",
         "14 bảng, 4 hàm, 1 trigger; chứa toàn bộ logic không gian"],
    ],
    widths=[1.1, 2.1, 2.9])

h3("Danh sách điểm cuối API")
table(
    ["Phương thức", "Đường dẫn", "Quyền", "Chức năng"],
    [
        ["POST", "/api/dang-nhap", "—", "Đăng nhập, trả về JWT"],
        ["GET", "/api/toi", "đã đăng nhập", "Thông tin nhân viên và khối MAP của họ"],
        ["GET", "/api/khu-vuc", "đã đăng nhập", "Danh sách khối MAP để vẽ bản đồ"],
        ["POST", "/api/cham-cong", "đã đăng nhập",
         "Kiểm tra bao hàm khối, áp R1–R4, ghi bản ghi và cảnh báo"],
        ["GET", "/api/cham-cong/lich-su", "đã đăng nhập", "Lịch sử của chính mình"],
        ["POST", "/api/giai-trinh", "đã đăng nhập", "Gửi đơn giải trình"],
        ["GET", "/api/giai-trinh", "quản lý", "Danh sách đơn"],
        ["PUT", "/api/giai-trinh/:id", "quản lý", "Duyệt hoặc từ chối đơn"],
        ["GET", "/api/canh-bao", "quản lý", "Danh sách cảnh báo bất thường"],
        ["GET", "/api/dashboard", "quản lý", "Số liệu tổng hợp"],
        ["GET", "/api/bao-cao/cong", "quản lý", "Báo cáo công theo tháng"],
    ],
    widths=[0.85, 1.6, 0.95, 2.7],
    caption="Bảng 3.2 — Các điểm cuối API của hệ thống")

h3("Giao diện đã cài đặt")
image(os.path.join(SHOT, "03-app-cham-cong.png"),
      "Hình 3.2 — Màn hình chấm công: thông tin làm việc lấy từ cơ sở dữ liệu, thanh "
      "chọn tầng, bản đồ ba chiều hiển thị các khối MAP kèm popup thuộc tính",
      folder=SHOT)
image(os.path.join(SHOT, "04-app-quan-ly.png"),
      "Hình 3.3 — Bảng điều khiển của quản lý: số liệu tổng hợp và danh sách cảnh báo "
      "bất thường sinh ra từ bốn quy tắc R1–R4", folder=SHOT)
page_break()

# ===========================================================================
#  CHƯƠNG 4
# ===========================================================================
h1("KẾT QUẢ VÀ ĐÁNH GIÁ")

h2("Kết quả bước đầu")
body_text(
    "Bản dựng thử nghiệm sử dụng dữ liệu thật: footprint toà nhà IFC One Saigon lấy từ "
    "OpenStreetMap, công trình 42 tầng cao 195,3 mét; cao độ nền giả định 5 mét nên "
    "chiều cao tầng trung bình là 4,65 mét. Ba văn phòng được đặt tại ba dải tầng khác "
    "nhau nhưng dùng chung một footprint, tức là có cùng toạ độ kinh độ và vĩ độ.")
table(
    ["Tình huống", "Cao độ điểm", "Khoảng cách ngang", "Kết quả hệ 3D", "Hàng rào 2D"],
    [
        ["Nhân viên Alpha Tech chấm công tại tầng 6", "28,25 m", "0 m", "HOP_LE", "Hợp lệ"],
        ["Nhân viên Alpha Tech chấm công tại tầng 22", "104,98 m", "0 m",
         "NGHI_NGO (lệch +62,78 m, tương đương 13,5 tầng)", "Hợp lệ — bỏ sót"],
        ["Nhân viên Alpha Tech chấm công cách toà nhà 270 m", "28,25 m", "270,2 m",
         "NGOAI_VUNG", "Ngoài vùng"],
    ],
    widths=[1.9, 0.85, 1.0, 1.5, 0.85],
    caption="Bảng 4.1 — Kết quả ba tình huống kiểm thử")
body_text(
    "Tình huống thứ hai là kết quả quan trọng nhất: với cùng một toạ độ ngang, hàng rào "
    "hai chiều kết luận hợp lệ trong khi mô hình MAP phát hiện được sai lệch 13,5 tầng "
    "và chuyển bản ghi sang trạng thái cần xác minh.")
image(os.path.join(SHOT, "02-ket-qua-nghi-ngo.png"),
      "Hình 4.1 — Kết quả tình huống chấm công sai tầng; điểm chấm công (hình cầu) nằm "
      "trong khối của doanh nghiệp khác", folder=SHOT)

para()
para("Kiểm chứng phần cơ sở dữ liệu", bold=True)
body_text(
    "Toàn bộ lược đồ và các hàm nghiệp vụ đã được chạy thật trên PostgreSQL 17.6 kết "
    "hợp PostGIS 3.6.2, không dừng ở mức mã nguồn. Kết quả trên cơ sở dữ liệu trùng "
    "khớp với bản cài đặt bằng JavaScript ở phía trình duyệt, xác nhận hai bên dùng "
    "chung một thuật toán.")
table(
    ["Hạng mục kiểm chứng", "Kết quả"],
    [
        ["Chạy schema.sql", "Thành công, tạo 14 bảng, 4 hàm và 1 trigger"],
        ["Chạy seed.sql", "Thành công, nạp dữ liệu toà nhà IFC One Saigon và ba văn phòng"],
        ["Trigger suy cao độ từ dải tầng",
         "Alpha 23,60–42,20 m · Beta 93,35–116,60 m · Gamma 163,10–191,00 m, "
         "đúng công thức ở mục 2.1.3"],
        ["Hàm kiem_tra_bao_ham()",
         "Đúng tầng → HOP_LE; sai tầng → NGHI_NGO; ngoài toà nhà → NGOAI_VUNG"],
        ["Hàm kiem_tra_r2() (dịch chuyển bất khả thi)",
         "Chấm công lại sau 3 phút ở vị trí cách 200 km → phát hiện bất thường; "
         "cùng toạ độ → không cảnh báo"],
    ],
    widths=[2.4, 3.7],
    caption="Bảng 4.2 — Kết quả chạy thật trên PostgreSQL 17.6 + PostGIS 3.6.2")

para()
para("Kiểm chứng toàn hệ thống", bold=True)
body_text(
    "Sau khi hoàn thiện máy chủ và giao diện, các luồng nghiệp vụ được chạy thử đầu "
    "cuối qua API thật trên cơ sở dữ liệu thật. Kết quả cho thấy bốn quy tắc phát hiện "
    "bất thường đều hoạt động, phân quyền chặn đúng, và luồng giải trình khép kín.")
table(
    ["Kịch bản kiểm thử", "Kết quả hệ thống trả về"],
    [
        ["Chấm công đúng tầng 6", "HOP_LE"],
        ["Khai báo tầng 22 trong khi thuê tầng 5–8",
         "NGHI_NGO, kèm cảnh báo R1: lệch 62,78 m tương đương 13,5 tầng"],
        ["Chấm công ở vị trí cách toà nhà 270 m",
         "NGOAI_VUNG, kèm cảnh báo R2 về dịch chuyển bất khả thi"],
        ["Thiết bị báo độ chính xác 0,3 m",
         "Cảnh báo R3: độ chính xác nhỏ bất thường, nghi giả lập vị trí"],
        ["Nhân viên khác dùng lại cùng một thiết bị",
         "Cảnh báo R4: thiết bị vừa được nhân viên khác dùng để chấm công"],
        ["Nhân viên gọi API dành cho quản lý",
         "HTTP 403 — không đủ quyền thực hiện chức năng này"],
        ["Quản lý duyệt đơn giải trình",
         "Bản ghi chuyển từ NGOAI_VUNG sang HOP_LE, các cảnh báo liên quan được đóng, "
         "báo cáo công cập nhật theo"],
    ],
    widths=[2.7, 3.4],
    caption="Bảng 4.3 — Kết quả kiểm thử đầu cuối trên hệ thống đã cài đặt")

h2("Ưu điểm")
bullets([
    "Giải quyết một hạn chế có thật của các hệ thống chấm công đang dùng phổ biến, "
    "không phải bài toán đặt ra chỉ để minh hoạ công nghệ.",
    "Trục cao độ là thành phần bắt buộc của lời giải, phù hợp trọng tâm môn học.",
    "Sinh ra trạng thái trung gian NGHI_NGO thay vì chỉ chấp nhận hoặc từ chối, phù hợp "
    "thực tế quản lý nhân sự vốn cần bước xác minh.",
    "Mô hình dữ liệu nhẹ, chỉ cần footprint và số tầng, có thể nhân rộng cho nhiều toà "
    "nhà mà không cần thiết bị khảo sát chuyên dụng.",
    "Trực quan hoá ba chiều tạo bằng chứng đối chiếu trực tiếp khi có tranh chấp về công.",
    "Bốn quy tắc phát hiện bất thường mở rộng giá trị của dữ liệu không gian sang bài "
    "toán chống gian lận.",
])

h2("Hạn chế")
bullets([
    "Độ cao đo bằng tín hiệu vệ tinh trong nhà rất kém; thuộc tính cao độ mà "
    "Geolocation API trả về thường rỗng hoặc có sai số hàng chục mét, lớn hơn cả chiều "
    "cao vài tầng. Hệ thống vì vậy phải chấp nhận nguồn cao độ hỗn hợp và mô hình chỉ "
    "đóng vai trò đối chiếu chứ không tự quyết định.",
    "Chưa sử dụng cảm biến bổ trợ như Wi-Fi RSSI, BLE beacon hay khí áp kế — vốn là "
    "giải pháp công nghiệp cho định vị trong nhà.",
    "Dữ liệu chiều cao công trình từ OpenStreetMap rất thiếu (khoảng 3,7 % số công "
    "trình có thuộc tính chiều cao), nên phần lớn phải nhập thủ công.",
    "Không thể chặn hoàn toàn việc giả lập vị trí ở mức hệ điều hành từ phía ứng dụng "
    "web; các quy tắc chỉ làm tăng chi phí gian lận.",
    "Mô hình ở mức LoD1 nên không phân biệt được các khu vực khác nhau trong cùng một tầng.",
])

h2("Hướng phát triển")
h3("Định vị trong nhà bằng BLE beacon")
body_text(
    "Bố trí thiết bị phát tín hiệu Bluetooth năng lượng thấp theo từng tầng và hợp nhất "
    "kết quả với mô hình MAP, cho phép xác định tầng độc lập với tín hiệu vệ tinh.")
h3("Nâng mức chi tiết lên LoD2 và LoD3")
body_text(
    "Mô tả từng phòng và khu vực làm việc theo chuẩn CityGML, phục vụ chấm công theo "
    "khu vực thay vì theo tầng, đồng thời hỗ trợ quản lý chỗ ngồi.")
h3("Hiệu chỉnh dung sai theo vùng che khuất")
body_text(
    "Phân tích vùng che khuất tín hiệu dựa trên mô hình khối của các toà nhà lân cận, "
    "từ đó tự động điều chỉnh dung sai ngang theo từng vị trí thay vì dùng một giá trị "
    "cố định cho toàn hệ thống.")
h3("Mở rộng cho công trường xây dựng")
body_text(
    "Khối chấm công có cao độ thay đổi theo tiến độ thi công, tức bổ sung chiều thứ tư "
    "là thời gian vào mô hình dữ liệu.")
h3("Tích hợp hệ thống nhân sự và tính lương")
body_text(
    "Kết nối dữ liệu công đã xác thực với hệ thống tính lương, rút ngắn quy trình đối "
    "soát cuối tháng.")

h2("Kết luận")
body_text(
    "Đề tài “Xây dựng hệ thống chấm công định vị 3D theo mô hình khối không gian phân "
    "tầng” cho thấy giá trị thực tiễn của việc bổ sung chiều thứ ba vào một hệ thống "
    "thông tin địa lý. Điểm cốt lõi không nằm ở việc hiển thị đẹp mắt các khối nhà, mà "
    "ở chỗ phép xác thực chấm công chỉ trở nên đúng khi mô hình dữ liệu có trục cao độ.")
body_text(
    "Mô hình MAP cho phép biểu diễn mỗi văn phòng bằng một khối không gian có giới hạn "
    "trên và giới hạn dưới, nhờ đó phân biệt được các doanh nghiệp thuê những tầng khác "
    "nhau trong cùng một cao ốc — điều mà hàng rào địa lý hai chiều không thể làm. Kết "
    "quả thử nghiệm đã chứng minh hệ thống bắt được tình huống chấm công sai tầng mà "
    "mô hình hai chiều ghi nhận là hợp lệ.")
body_text(
    "Đề tài đồng thời thể hiện khả năng kết hợp giữa hệ thống thông tin địa lý ba "
    "chiều, cơ sở dữ liệu không gian, công nghệ web và nghiệp vụ quản lý nhân sự; đồng "
    "thời tạo nền tảng để phát triển thành một hệ thống định vị trong nhà hoàn chỉnh "
    "trong tương lai.")
page_break()

# ===========================================================================
#  TÀI LIỆU THAM KHẢO
# ===========================================================================
h1("TÀI LIỆU THAM KHẢO", numbered=False)
for line in [
    "[1] Esri (2024). ArcGIS Maps SDK for JavaScript – SceneView, "
    "https://developers.arcgis.com/javascript/latest/api-reference/esri-views-SceneView.html",
    "[2] Esri (2024). ExtrudeSymbol3DLayer, "
    "https://developers.arcgis.com/javascript/latest/api-reference/"
    "esri-symbols-ExtrudeSymbol3DLayer.html",
    "[3] PostGIS Development Group (2024). PostGIS Documentation – Geography Type and "
    "ST_DWithin, https://postgis.net/docs/",
    "[4] Open Geospatial Consortium (2012). OGC City Geography Markup Language "
    "(CityGML) Encoding Standard, https://www.ogc.org/standards/citygml",
    "[5] OpenStreetMap contributors (2026). Overpass API, "
    "https://wiki.openstreetmap.org/wiki/Overpass_API",
    "[6] OpenStreetMap Wiki (2026). Simple 3D Buildings, "
    "https://wiki.openstreetmap.org/wiki/Simple_3D_buildings",
    "[7] W3C (2022). Geolocation API Specification, https://www.w3.org/TR/geolocation/",
    "[8] MDN Web Docs (2026). GeolocationCoordinates.altitude, "
    "https://developer.mozilla.org/en-US/docs/Web/API/GeolocationCoordinates/altitude",
    "[9] Phan Thanh Vũ (2026). Bài giảng ArcGIS JS – Lab 1 và Lab 2, môn Hệ thống "
    "thông tin địa lý 3 chiều, Trường Đại học Công nghệ Thông tin – ĐHQG TP.HCM.",
]:
    p = para(line)
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.first_line_indent = Inches(-0.35)

# ===========================================================================
#  PHỤ LỤC
# ===========================================================================
h1("PHỤ LỤC A. DANH SÁCH CHỨC NĂNG CHO DEMO", numbered=False)
bullets([
    "Đăng nhập với ba vai trò khác nhau.",
    "Chấm công bằng GPS thật trên điện thoại qua kết nối HTTPS.",
    "Chấm công mô phỏng bằng cách chọn vị trí và tầng trên bản đồ ba chiều.",
    "Trình diễn ba tình huống: đúng tầng, sai tầng và ngoài toà nhà.",
    "So sánh trực tiếp kết luận của mô hình ba chiều với hàng rào hai chiều.",
    "Thay đổi dung sai ngang và dung sai đứng, quan sát kết quả thay đổi theo.",
    "Dựng một khối vùng chấm công mới trực tiếp trên bản đồ.",
    "Xem danh sách cảnh báo bất thường và mở lại vị trí ba chiều của bản ghi.",
    "Gửi và duyệt một đơn giải trình.",
    "Xuất báo cáo công của một phòng ban.",
])

h1("PHỤ LỤC B. KẾ HOẠCH TRIỂN KHAI THEO GIAI ĐOẠN", numbered=False)
table(
    ["Giai đoạn", "Nội dung công việc", "Sản phẩm bàn giao"],
    [
        ["Giai đoạn 1", "Khảo sát toà nhà, thu thập footprint, số tầng và chiều cao; "
                        "chuẩn hoá dữ liệu về GeoJSON",
         "Bộ dữ liệu không gian của toà nhà thử nghiệm"],
        ["Giai đoạn 2", "Thiết kế và cài đặt cơ sở dữ liệu PostGIS; viết hàm kiểm tra "
                        "bao hàm khối và hàm phát hiện bất thường",
         "Lược đồ cơ sở dữ liệu, script khởi tạo và dữ liệu mẫu"],
        ["Giai đoạn 3", "Dựng bản đồ ba chiều: khối toà nhà, khối MAP, điểm chấm công",
         "Màn hình bản đồ ba chiều chạy được"],
        ["Giai đoạn 4", "Xây dựng API và giao diện chấm công; tích hợp Geolocation API",
         "Luồng chấm công hoàn chỉnh trên HTTPS"],
        ["Giai đoạn 5", "Xây dựng luồng giải trình, duyệt và bảng điều khiển quản trị",
         "Giao diện quản lý và báo cáo công"],
        ["Giai đoạn 6", "Kiểm thử thực địa ba tình huống; đo thời gian phản hồi; hoàn "
                        "thiện báo cáo",
         "Bảng kết quả kiểm thử, video trình diễn, báo cáo cuối kỳ"],
    ],
    widths=[0.9, 2.9, 2.2],
    caption="Bảng B.1 — Kế hoạch triển khai theo giai đoạn")

# ---------------------------------------------------------------------------
# --- dong bo font: Times New Roman, than bai 13pt (giong file mau) ---
def _ep_font(paragraphs):
    for p in paragraphs:
        for r in p.runs:
            if r.font.name == "Consolas":
                continue
            r.font.name = "Times New Roman"
            rPr = r._element.get_or_add_rPr()
            rf = rPr.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts")
                rPr.insert(0, rf)
            for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                rf.set(qn(attr), "Times New Roman")
            if r.font.size is None and not p.style.name.startswith("Heading"):
                r.font.size = Pt(13)


_ep_font(doc.paragraphs)
for _t in doc.tables:
    for _row in _t.rows:
        for _cell in _row.cells:
            _ep_font(_cell.paragraphs)

doc.save(OUT_DOCX)
print("DONE ->", OUT_DOCX)
